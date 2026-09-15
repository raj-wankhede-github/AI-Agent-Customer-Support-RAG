import io
import zipfile
from pathlib import Path

import docx
import pytest

from app.core.errors import IngestionError, PayloadTooLargeError, UnsupportedMediaTypeError, ValidationAppError
from app.ingestion.loaders.docx import DocxLoader
from app.ingestion.loaders.html import HtmlLoader
from app.ingestion.loaders.pdf import PdfLoader
from app.ingestion.loaders.text import MarkdownLoader, TextLoader
from app.ingestion.validation import validate_upload
from app.models.enums import SourceType

SEED = Path(__file__).resolve().parents[2] / "seed" / "acme"


def test_markdown_front_matter_headings_lists_and_tables():
    data = (
        b"---\ntitle: Refunds\n---\n# Refunds\n\nIntro **text** here.\n\n## Window\n\n- Item one\n- Item two\n\n"
        b"| A | B |\n|---|---|\n| 1 | 2 |\n"
    )
    document = MarkdownLoader().load(data, "refunds.md")
    assert [b.kind for b in document.blocks] == ["heading", "paragraph", "heading", "list", "table"]
    assert document.title == "Refunds"
    assert document.blocks[1].text == "Intro text here."
    assert document.blocks[4].text == "A | B\n1 | 2"


def test_html_ignores_scripts_navigation_and_footer():
    html = b"""<html><head><title>FAQ</title><link rel="canonical" href="https://help.example/faq">
    <script>var secret = "analytics";</script></head><body><nav>Home | Orders</nav>
    <main><h1>Account FAQ</h1><p>Reset your password from the sign-in page.</p>
    <ul><li>Step one</li><li>Step two</li></ul><table><tr><th>Plan</th><th>Price</th></tr><tr><td>Pro</td><td>$10</td></tr></table>
    </main><footer>Copyright</footer></body></html>"""
    document = HtmlLoader().load(html, "faq.html")
    text = "\n".join(b.text for b in document.blocks)
    assert "analytics" not in text and "Home | Orders" not in text and "Copyright" not in text
    assert [b.kind for b in document.blocks] == ["heading", "paragraph", "list", "table"]
    assert document.title == "Account FAQ"
    assert document.metadata["source_uri"] == "https://help.example/faq"


def test_seed_pdf_keeps_pages_and_headings():
    document = PdfLoader(max_pages=500).load((SEED / "shipping-policy.pdf").read_bytes(), "shipping-policy.pdf")
    assert document.page_count == 2
    headings = {b.text: b.page for b in document.blocks if b.kind == "heading"}
    assert headings["Standard Shipping"] == 1
    assert headings["International Shipping"] == 2
    assert any("3-5 business days" in b.text for b in document.blocks if b.kind == "paragraph")


def test_corrupt_pdf_fails_with_classified_error():
    with pytest.raises(IngestionError) as exc:
        PdfLoader(max_pages=500).load(b"%PDF-1.4\nthis is not really a pdf", "broken.pdf")
    assert exc.value.error_code == "CORRUPT_FILE"


def test_docx_headings_lists_and_tables():
    source = docx.Document()
    source.add_heading("Guide", level=1)
    source.add_paragraph("Welcome to the guide.")
    source.add_paragraph("First step", style="List Bullet")
    table = source.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Weight", "180 g"
    buffer = io.BytesIO()
    source.save(buffer)
    document = DocxLoader().load(buffer.getvalue(), "guide.docx")
    assert [b.kind for b in document.blocks] == ["heading", "paragraph", "list", "table"]
    assert document.blocks[3].text == "Weight | 180 g"


def test_plain_text_detects_uppercase_headings():
    document = TextLoader().load((SEED / "support-contact-policy.txt").read_bytes(), "support.txt")
    headings = [b.text for b in document.blocks if b.kind == "heading"]
    assert "SUPPORT HOURS" in headings and "CONTACT CHANNELS" in headings


def test_upload_validation_accepts_supported_types():
    assert validate_upload("notes.md", b"# Hello", "text/markdown", 1000).source_type is SourceType.MARKDOWN
    real_docx = (SEED / "smarthub-product-guide.docx").read_bytes()
    assert validate_upload("guide.docx", real_docx, None, 10_000_000).source_type is SourceType.DOCX


@pytest.mark.parametrize(
    ("filename", "data", "content_type", "error"),
    [
        ("tool.exe", b"MZ....", None, UnsupportedMediaTypeError),
        ("empty.txt", b"", None, ValidationAppError),
        ("fake.pdf", b"hello world", None, UnsupportedMediaTypeError),
        ("binary.txt", b"abc\x00def", None, UnsupportedMediaTypeError),
        ("mismatch.pdf", b"%PDF-1.4 ...", "text/html", UnsupportedMediaTypeError),
        ("latin1.txt", "caf\xe9".encode("latin-1"), None, ValidationAppError),
        (
            "active.pdf",
            b"%PDF-1.4\n1 0 obj << /OpenAction << /S /JavaScript /JS (app.alert(1)) >> >>",
            None,
            ValidationAppError,
        ),
    ],
)
def test_upload_validation_rejects_unsafe_or_invalid_files(filename, data, content_type, error):
    with pytest.raises(error):
        validate_upload(filename, data, content_type, 1000)


def test_upload_validation_enforces_size_limit():
    with pytest.raises(PayloadTooLargeError):
        validate_upload("big.txt", b"a" * 2000, None, 1000)


def test_macro_enabled_docx_is_rejected():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            "<Types><Override ContentType='application/vnd.ms-word.document.macroEnabled.main+xml'/></Types>",
        )
        archive.writestr("word/document.xml", "<document/>")
        archive.writestr("word/vbaProject.bin", b"\x00\x01")
    with pytest.raises(ValidationAppError):
        validate_upload("macro.docx", buffer.getvalue(), None, 1_000_000)
