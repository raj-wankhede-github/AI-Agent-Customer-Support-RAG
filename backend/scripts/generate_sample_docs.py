"""Generates the binary seed documents (PDF, DOCX) from the text below.

The committed files in seed/acme/ were produced by this script; re-run it after editing:
    uv run python scripts/generate_sample_docs.py

All content is fictional.
"""

from __future__ import annotations

import textwrap
from pathlib import Path

import docx

ROOT = Path(__file__).resolve().parent.parent
ACME = ROOT / "seed" / "acme"

SHIPPING_PAGES = [
    [
        "# Acme Shipping Policy",
        "Effective date: 2026-01-01",
        "This policy explains how Acme ships orders placed on acme.example.",
        "# Order Processing",
        "Orders placed before 2:00 p.m. Eastern Time on a business day are processed the same day. Orders placed "
        "after 2:00 p.m. Eastern Time, or on weekends and public holidays, are processed on the next business day.",
        "# Standard Shipping",
        "Standard shipping takes 3-5 business days after your order ships. Standard shipping is free for orders of "
        "$50 or more. Orders under $50 pay a flat standard shipping fee of $5.99.",
        "# Expedited Shipping",
        "Expedited shipping takes 1-2 business days after your order ships and costs $14.99. Expedited shipping is "
        "available to addresses in the contiguous United States only.",
    ],
    [
        "# International Shipping",
        "International shipping is available to Canada, the United Kingdom, the European Union and Australia. "
        "International orders typically arrive within 7-14 business days after they ship. Customs duties and "
        "import taxes are the responsibility of the recipient and are not included in the shipping fee.",
        "# Order Tracking",
        "A tracking number is emailed to you when your order ships. Tracking information can take up to 24 hours "
        "to appear on the carrier website.",
        "# Lost or Damaged Packages",
        "If your package has not arrived within 10 business days after the estimated delivery date, or arrives "
        "damaged, contact Acme Support within 30 days of the estimated delivery date so we can open a carrier claim.",
    ],
]


def _pdf_escape(value: str) -> str:
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def write_pdf(path: Path, pages: list[list[str]], title: str) -> None:
    """Minimal text-layer PDF (Helvetica, WinAnsi) - enough for real text extraction."""
    objects: list[bytes] = []

    def add(obj: str | bytes) -> int:
        objects.append(obj.encode("latin-1") if isinstance(obj, str) else obj)
        return len(objects)

    catalog = add("")  # placeholder
    pages_obj = add("")
    font = add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    bold = add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>")
    page_ids = []
    for lines in pages:
        ops = ["BT", "15 TL", "72 740 Td"]
        for item in lines:
            if item.startswith("# "):
                ops += ["T*", "/F2 14 Tf", f"({_pdf_escape(item[2:])}) Tj", "T*", "/F1 11 Tf"]
            else:
                for wrapped in textwrap.wrap(item, 92):
                    ops += ["/F1 11 Tf", f"({_pdf_escape(wrapped)}) Tj", "T*"]
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        content = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        page_ids.append(add(
            f"<< /Type /Page /Parent {pages_obj} 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font} 0 R /F2 {bold} 0 R >> >> /Contents {content} 0 R >>"
        ))  # fmt: skip
    objects[catalog - 1] = f"<< /Type /Catalog /Pages {pages_obj} 0 R >>".encode()
    kids = " ".join(f"{p} 0 R" for p in page_ids)
    objects[pages_obj - 1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()
    info = add(f"<< /Title ({_pdf_escape(title)}) /Producer (Acme sample generator) >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root {catalog} 0 R /Info {info} 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


def write_product_guide(path: Path) -> None:
    document = docx.Document()
    document.core_properties.title = "Acme SmartHub Product Guide"
    document.add_heading("Acme SmartHub Product Guide", level=1)
    document.add_paragraph(
        "The Acme SmartHub connects your Acme smart lights, plugs and sensors so you can control them from the "
        "Acme Home app."
    )
    document.add_heading("Technical Specifications", level=2)
    specs = [
        ("Specification", "Value"),
        ("Wi-Fi", "802.11ax (Wi-Fi 6), dual-band 2.4 GHz and 5 GHz"),
        ("Bluetooth", "Bluetooth 5.3"),
        ("Zigbee", "Zigbee 3.0, up to 64 connected devices"),
        ("Power", "5 V / 2 A via USB-C (adapter included)"),
        ("Dimensions", "98 x 98 x 32 mm"),
        ("Weight", "180 g"),
        ("Operating temperature", "0 to 40 degrees Celsius"),
    ]
    table = document.add_table(rows=0, cols=2)
    for key, value in specs:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = key, value
    document.add_heading("Setting Up Your SmartHub", level=2)
    for step in (
        "Plug the SmartHub into power using the included USB-C adapter.",
        "Download the Acme Home app and sign in or create an Acme account.",
        "Press and hold the pairing button on the back of the SmartHub for 5 seconds until the light blinks blue.",
        "In the Acme Home app, tap Add Device and follow the on-screen instructions.",
    ):
        document.add_paragraph(step, style="List Number")
    document.add_heading("Status Light Meanings", level=2)
    for meaning in (
        "Solid green: the SmartHub is connected and working normally.",
        "Blinking blue: the SmartHub is in pairing mode.",
        "Solid red: the SmartHub cannot reach the internet.",
        "Blinking amber: a firmware update is being installed. Do not unplug the SmartHub.",
    ):
        document.add_paragraph(meaning, style="List Bullet")
    document.add_heading("Factory Reset", level=2)
    document.add_paragraph(
        "To factory reset the SmartHub, press and hold the reset button for 15 seconds until the status light "
        "flashes red. A factory reset removes all paired devices and settings."
    )
    document.add_heading("Warranty", level=2)
    document.add_paragraph(
        "The Acme SmartHub includes a 2-year limited warranty from the date of purchase. The warranty covers "
        "manufacturing defects and does not cover accidental damage, water damage or unauthorized modifications."
    )
    document.save(str(path))


if __name__ == "__main__":
    ACME.mkdir(parents=True, exist_ok=True)
    write_pdf(ACME / "shipping-policy.pdf", SHIPPING_PAGES, "Acme Shipping Policy")
    write_product_guide(ACME / "smarthub-product-guide.docx")
    print("wrote", ACME / "shipping-policy.pdf", "and", ACME / "smarthub-product-guide.docx")
